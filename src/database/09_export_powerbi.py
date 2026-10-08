"""Export the RFI star schema so Power BI can load it without DuckDB.

Reads data/processed/rfi.duckdb and does not change it. Writes
data/processed/powerbi/rfi_star_schema.xlsx (one sheet per table) and one
CSV per table in the same folder. Each run replaces those files.
"""

from __future__ import annotations

from pathlib import Path

try:
    import duckdb
except ImportError as error:
    raise SystemExit(
        "Install duckdb from requirements.txt before running this script."
    ) from error

try:
    from openpyxl import Workbook
except ImportError as error:
    raise SystemExit(
        "Install openpyxl from requirements.txt before running this script."
    ) from error

ROOT = Path(__file__).resolve().parents[2]
DATABASE = ROOT / "data" / "processed" / "rfi.duckdb"
OUT_DIR = ROOT / "data" / "processed" / "powerbi"
WORKBOOK = OUT_DIR / "rfi_star_schema.xlsx"
TABLE_ORDER = (
    "dim_date",
    "dim_party",
    "dim_discipline",
    "dim_issue_type",
    "fact_rfi",
)
SORT_COLUMN = {
    "dim_date": "date_key",
    "dim_party": "party_key",
    "dim_discipline": "discipline_key",
    "dim_issue_type": "issue_type_key",
    "fact_rfi": "rfi_key",
}


def print_progress(done: int, total: int, label: str) -> None:
    percent = 100 if total == 0 else done * 100 // total
    print(f"\r{label}: {percent}% ({done}/{total})", end="", flush=True)


def table_names(connection: duckdb.DuckDBPyConnection) -> list[str]:
    rows = connection.execute(
        """
        SELECT table_name
        FROM information_schema.tables
        WHERE table_schema = 'main'
          AND table_type = 'BASE TABLE'
        """
    ).fetchall()
    found = {name for (name,) in rows}
    ordered = [name for name in TABLE_ORDER if name in found]
    ordered.extend(sorted(found - set(ordered)))
    return ordered


def sql_path(path: Path) -> str:
    return str(path).replace("'", "''")


def export_csv(connection: duckdb.DuckDBPyConnection, name: str) -> int:
    sort = SORT_COLUMN.get(name, "1")
    query = f'SELECT * FROM "{name}" ORDER BY {sort}'
    csv_path = sql_path(OUT_DIR / f"{name}.csv")
    connection.execute(
        f"COPY ({query}) TO '{csv_path}' (HEADER, DELIMITER ',')"
    )
    count = connection.execute(f'SELECT count(*) FROM "{name}"').fetchone()
    return int(count[0])


def write_workbook(connection: duckdb.DuckDBPyConnection, names: list[str]) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name in names:
        sheet = workbook.create_sheet(name[:31])
        sort = SORT_COLUMN.get(name, "1")
        result = connection.execute(f'SELECT * FROM "{name}" ORDER BY {sort}')
        sheet.append([column[0] for column in result.description])
        for row in result.fetchall():
            sheet.append(list(row))
    workbook.save(WORKBOOK)


def main() -> None:
    if not DATABASE.exists():
        raise SystemExit(
            "Missing data/processed/rfi.duckdb. Run src/database/07_load_rfi_duckdb.py first."
        )

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    if WORKBOOK.exists():
        WORKBOOK.unlink()

    connection = duckdb.connect(str(DATABASE), read_only=True)
    names = table_names(connection)
    if not names:
        connection.close()
        raise SystemExit("The database has no tables.")

    counts: list[tuple[str, int]] = []
    for index, name in enumerate(names, start=1):
        print_progress(index, len(names), "Exporting tables")
        counts.append((name, export_csv(connection, name)))
    write_workbook(connection, names)
    connection.close()
    print()

    print(f"Wrote {WORKBOOK.relative_to(ROOT)}")
    for name, count in counts:
        print(f"  {count:3} {name}.csv")
    print("In Power BI: Get data > Excel workbook > rfi_star_schema.xlsx")


if __name__ == "__main__":
    main()
